"""DOCX 报告导出 —— 用 python-docx,纯 Python 实现,无平台限制。"""
import io
from datetime import datetime

from docx import Document


def build_health_report_docx(start, end, algorithm, rows):
    doc = Document()
    doc.add_heading("健康度历史报告", level=1)
    doc.add_paragraph(f"时间范围：{start or '--'} 至 {end or '--'}")
    doc.add_paragraph(f"评估算法：{algorithm or '--'}")
    doc.add_paragraph(f"生成时间：{datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")

    table = doc.add_table(rows=1, cols=5)
    table.style = "Light Grid Accent 1"
    hdr = table.rows[0].cells
    for i, name in enumerate(["ID", "记录时间", "模式", "电机1温度", "电机2温度"]):
        hdr[i].text = name
    for row in rows:
        cells = table.add_row().cells
        cells[0].text = str(row.get("id", ""))
        cells[1].text = str(row.get("timestamp", ""))
        cells[2].text = str(row.get("sys_mode", ""))
        cells[3].text = str(row.get("m1_temp", ""))
        cells[4].text = str(row.get("m2_temp", ""))

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf


def build_fault_report_docx(fault_row):
    doc = Document()
    doc.add_heading("故障详情报告", level=1)
    report = fault_row.get("fault_report", {})
    doc.add_paragraph(f"事件ID：{report.get('event_id', fault_row.get('id', ''))}")
    doc.add_paragraph(f"故障时间：{fault_row.get('timestamp', '')}")
    doc.add_paragraph(f"故障等级：{fault_row.get('error_type', '')}")
    doc.add_paragraph(f"故障设备：{report.get('故障设备', '--')}")
    doc.add_paragraph(f"故障部件：{report.get('故障部件', '--')}")
    doc.add_paragraph(f"故障类型：{report.get('故障类型', '--')}")

    abnormal = report.get("异常参数", [])
    if abnormal:
        doc.add_heading("异常参数", level=2)
        table = doc.add_table(rows=1, cols=4)
        table.style = "Light Grid Accent 1"
        hdr = table.rows[0].cells
        for i, name in enumerate(["参数名称", "当前值", "阈值", "说明"]):
            hdr[i].text = name
        for item in abnormal:
            cells = table.add_row().cells
            cells[0].text = str(item.get("name", ""))
            cells[1].text = str(item.get("current", ""))
            cells[2].text = str(item.get("threshold", ""))
            cells[3].text = str(item.get("desc", ""))

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf
